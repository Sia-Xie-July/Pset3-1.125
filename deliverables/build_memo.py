from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
r=json.loads((root/'deliverables/financial-results.json').read_text())
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.62);sec.left_margin=sec.right_margin=Inches(.7)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(5)
for n in ['Title','Heading 1','Heading 2']:
 st=doc.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(5);st.paragraph_format.space_before=Pt(8)
doc.styles['Title'].font.size=Pt(21);doc.styles['Heading 1'].font.size=Pt(12)
text=[]
def p(t):doc.add_paragraph(t);text.append(t)
def h(t):doc.add_heading(t,1);text.append('\n## '+t)
doc.add_heading('Investment recommendation for shared university AI computing',0)
p('University Consortium Investment Committee | 6 October 2026 | Planning study in constant 2026 EUR')
h('Recommendation')
p('Approve a capped diligence and leased-compute pilot in Finland; return the 25 MW construction proposal for more evidence. Retain Kajaani and a conditional 10 MW-facility hybrid phase as the working hypothesis. The hybrid has the lowest modeled base-case cost, but leasing wins when demand halves. No institution has signed demand commitments and no utility offer or equipment quote has been established.')
h('Demand and options')
p('Retain 20 MW IT × PUE 1.25 = 25 MW total facility load as the full-build comparison. A 1.4 kW/GPU-equivalent IT allowance implies 14,280 GPUs. At 97% scheduling availability and assumed 25% productive utilization, demand is 30.34 million useful GPU-hours/year. This is an illustrative forecast, not surveyed demand. The hybrid installs 5,712 GPU equivalents (about 8 MW IT / 10 MW facility).')
p('All options serve the same demand over years 1–10. Full build opens in year 3; hybrid opens in year 4; both lease capacity before opening and for overflow. Grid delay moves owned opening one year later. Half utilization reduces demand by half while retaining idle power and fixed costs. GPU purchases wait until the year before actual opening; replacement occurs after four operating years.')
h('Required financial comparison')
t=doc.add_table(rows=1,cols=6);t.autofit=False
heads=['Case / option','Pre-open €m','Annual opex €m','€/useful hour','Capital risk €m','Cost NPV €m']
for c,v in zip(t.rows[0].cells,heads):c.text=v
for x in r['results']:
 cells=t.add_row().cells
 vals=[x['scenario']+' / '+x['option'],f"{x['preOpeningCash']/1e6:.1f}",f"{x['annualOperatingCost']/1e6:.1f}",f"{x['costPerGPUHour']:.2f}",f"{x['capitalAtRisk']/1e6:.1f}",f"{x['npvCost']/1e6:.1f}"]
 for c,v in zip(cells,vals):c.text=v
for rowno,row in enumerate(t.rows):
 for i,c in enumerate(row.cells):
  c.width=Inches(1.36 if i==0 else 1.09)
  tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
  tcPr.append(b)
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF5' if rowno==0 else 'FFFFFF');tcPr.append(sh)
  for para in c.paragraphs:
   para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
   for run in para.runs:run.font.size=Pt(9);run.bold=rowno==0
p('Opex is the first owned-operation year, or year 1 for lease. Useful-hour cost discounts resource costs and delivered hours at 6%. Pre-opening cash includes capital, bridge operating costs and interest, before debt draw. Capital risk is unrecoverable initial capital plus pre-opening operating costs/interest and three-month cloud cancellation exposure; it is not added again to NPV.')
h('Assumptions and limitations')
p('Unquoted allowances: facility €10m/IT MW; grid €15m full build (proportional for hybrid); installed GPU/host/fabric/storage €40,000 per GPU; electricity €0.10/kWh; idle IT 35% of nameplate. Staffing €3m/year at full scale, facility maintenance/insurance 2% and IT support 3% of capex. Cloud: observed $3.99/GPU-hour [S14], assumed €0.92/$, 85% useful/billed efficiency and 10% storage/network uplift. Large-cluster availability and workload equivalence are unverified; test the published $5.54 cluster rate as well.')
doc.add_page_break()
h('Technical concept and failure response')
p('Kajaani uses grid power, direct-to-chip cooling for compatible servers, closed loops, dry coolers and supplemental refrigeration [S1, S3]. Two independent 25 MW electrical paths are proposed: losing one must leave 25 MW usable. Six site-rated 5 MW generators leave 25 MW after one fails. Ten-minute UPS energy is 4.17 MWh delivered, or at least 5.79 MWh nameplate assuming 90% efficiency and 80% end-of-life capacity. For 48 hours at 25 MW, 1,200 MWh requires an assumed 345,000 litres diesel at 0.25 L/kWh plus 15% reserve. None of these assumptions proves uptime; verify topology, switching, fuel, cooling continuity, permits and integrated tests. A second generator failure requires at least 5 MW load shedding, prioritizing essential services over training.')
p('Propose six 5 MW cooling trains with one spare, two diverse 100 Gbit/s external links, workload-sized internal fabric and 5 PB usable storage. Water demand, actual PUE, tariff, grid date and site capacity remain unknown. No solar/wind energy, heat revenue or zero-water benefit is credited. Hourly energy and site water studies are required; annual national electricity shares do not establish site emissions.')
h('Financing and governance')
p('Diligence equity precedes binding commitments. Construction debt requires funded member reservations, land/permits, a grid offer, bounded EPC cost and independent engineering review. Equipment finance follows energization milestones and acceptance benchmarks. Illustrative finance is 50% initial-capex debt, 4% real interest, eight operating-year principal installments; replacement is member-funded. The online model shows interest, equity cash and outstanding year-10 debt separately. Resource NPV excludes financing and credits no terminal resale, tax shields or grants; these are conservative conventions, not unknown measurements entered as zero.')
p('If approved, a consortium nonprofit owns the facility and shared IT, contracting construction, operations, power and support. One institution one board vote; publish reservation fees and metered variable charges. Allocate 60% research, 25% teaching, 15% inference initially; protect small-institution access and cap routine reservations by a single member at 30% unless openly approved. Release unused reservations to a common pool. Members fund their committed fixed costs; exit payments cover noncancelable obligations. Equity bears uninsured delay/shortfall; credit utility/EPC compensation only when contracted.')
h('Three findings most likely to reverse the decision')
p('1. Contracted demand and alternatives. A pilot and signed funded commitments establish useful hours, concurrency and service levels beyond what existing academic or commercial facilities can supply. CSC offers eligible academic LUMI access [S17], but sufficient allocation and software fit are unverified. At half demand, modeled leasing costs €4.79/useful hour versus €6.17 for hybrid.')
p('2. Utility and engineering feasibility. A binding offer establishes connection capacity, cost and date; equipment and integrated failure studies establish cooling, water, fiber diversity and backup capability. A delayed hybrid increases modeled cost to €4.35/useful hour. A shared failure path, costly upgrades or permitting constraint can favor lease or relocation.')
p('3. Comparable delivered-compute prices. Competitive quotes, benchmarked productivity, cloud availability and GPU renewal prices change the ranking. Compare Canada/Québec and Singapore as well as Kajaani: annual electricity statistics differ in scope, Québec projects ≥5 MW require authorization [S20], and Singapore’s hot humid climate affects cooling [S19]. Do not select a country from an unsupported composite score.')
h('Evidence and approval request')
p('Approve only a capped 8–12 week pilot and diligence budget agreed by members. No full build or phase procurement until the three gates above pass. Sources and complete assumptions are linked in the website Evidence and Decision dossier pages. Reference numbers identify D1 source records; exact site facts and quotations remain unresolved.')
p('[S1] LUMI FAQ, lumi-supercomputer.eu/faq/ · [S3] DOE/FEMP data center design guide (2024), energy.gov · [S14] Lambda pricing, lambda.ai/pricing · [S17] CSC LUMI service, research.csc.fi/service/lumi-supercomputer/ · [S19] Singapore climate, weather.gov.sg/climate-climate-of-singapore/ · [S20] Québec connection authorization, quebec.ca. Public sources reviewed 5–6 October 2026.')
for node in list(doc.element.xpath('//w:pBdr')) + list(doc.styles.element.xpath('//w:pBdr')):
 node.getparent().remove(node)
doc.save(root/'investment-memo.docx')
(root/'investment-memo.md').write_text('# Investment recommendation for shared university AI computing\n\n'+'\n\n'.join(text)+'\n\nNumerical scenario rows and full annual cash flows: [financial-results.json](deliverables/financial-results.json).\n')
print('Created memo')
